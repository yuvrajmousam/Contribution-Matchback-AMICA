import streamlit as st
import pandas as pd
import numpy as np
import io
import time
from datetime import date
from openpyxl import load_workbook
from openpyxl.styles import numbers, PatternFill, Font

# =========================
# APP CONFIGURATION
# =========================
st.set_page_config(
    page_title="Granular Spec Multiplier",
    layout="wide",
    initial_sidebar_state="collapsed",
)
st.title("📊 Contribution Matchback of BOSE-US")

# =========================
# CUSTOM CSS: JUGGLING BOSE LOADER
# =========================
st.markdown("""
<style>
    /* The Overlay (Background) */
    .bose-loader-overlay {
        position: fixed;
        top: 0;
        left: 0;
        width: 100vw;
        height: 100vh;
        background: rgba(255, 255, 255, 0.95);
        z-index: 999999;
        display: flex;
        justify-content: center;
        align-items: center;
        flex-direction: column;
        backdrop-filter: blur(4px);
    }

    /* Container for the letters */
    .bose-juggler {
        display: flex;
        justify-content: center;
        align-items: flex-end;
        gap: 15px;
        height: 100px;
        margin-bottom: 20px;
    }

    /* Individual Letters */
    .bose-letter {
        font-family: 'Helvetica Neue', Helvetica, Arial, sans-serif;
        font-weight: 900;
        font-size: 80px;
        color: #000000;
        line-height: 1;
        animation: juggle 1.4s ease-in-out infinite;
    }

    /* Staggered Delay for the "Wave/Juggle" Effect */
    .bose-letter:nth-child(1) { animation-delay: 0.0s; }
    .bose-letter:nth-child(2) { animation-delay: 0.15s; }
    .bose-letter:nth-child(3) { animation-delay: 0.3s; }
    .bose-letter:nth-child(4) { animation-delay: 0.45s; }

    /* Status Text */
    .bose-status {
        color: #333; 
        font-family: sans-serif;
        font-size: 18px;
        font-weight: 600;
        letter-spacing: 2px;
        text-transform: uppercase;
        animation: fade 1.5s ease-in-out infinite alternate;
    }

    /* Keyframes */
    @keyframes juggle {
        0%, 100% { transform: translateY(0); }
        50% { transform: translateY(-40px); }
    }
    @keyframes fade {
        from { opacity: 0.6; }
        to { opacity: 1; }
    }
</style>
""", unsafe_allow_html=True)

# =========================
# LOADING HELPERS
# =========================

@st.cache_data
def load_main_spec(file_like_object):
    try:
        xl = pd.ExcelFile(file_like_object)
        sheet_name = next((s for s in xl.sheet_names if "model" in s.lower()), None)
        if not sheet_name:
            st.error("❌ No sheet containing 'Model' found in Main Spec.")
            return None

        preview = pd.read_excel(file_like_object, sheet_name=sheet_name, nrows=50, header=None)
        header_row = None

        for i, row in preview.iterrows():
            cells = [str(x).strip().lower() for x in row.fillna("")]
            short_texts = [c for c in cells if 0 < len(c) <= 20]
            if len(short_texts) < 2:
                continue
            if any("variable" in c for c in short_texts) and any("type" in c for c in short_texts):
                header_row = i
                break
        
        if header_row is None:
            st.error("❌ Could not auto-detect header row (with 'Variable' and 'Type') in Main Spec.")
            return None

        df = pd.read_excel(file_like_object, sheet_name=sheet_name, header=header_row)
        df.columns = [str(c).strip().upper() for c in df.columns]

        possible_var_cols = ["VARIABLE", "VARIABLES", "VARIABLE NAME", "VAR NAME", "VARIABLE_NAME"]
        possible_type_cols = ["TYPE", "TYPES", "VARIABLE TYPE"]

        var_col = next((c for c in df.columns if c in possible_var_cols), None)
        type_col = next((c for c in df.columns if c in possible_type_cols), None)
        if not var_col or not type_col:
            st.error(f"❌ 'Variable' or 'Type' column not found in Main Spec. Found: {df.columns.tolist()}")
            return None

        df = df[[var_col, type_col]].rename(columns={var_col: "VARIABLE", type_col: "TYPE"})
        df["VARIABLE"] = df["VARIABLE"].astype(str).str.strip().str.upper()
        df["TYPE"] = df["TYPE"].astype(str).str.strip().str.title()
        return df
    except Exception as e:
        st.error(f"Error loading Main Spec: {e}")
        return None

@st.cache_data
def load_pmf(file_like_object):
    try:
        xl = pd.ExcelFile(file_like_object)
        sheet_name = next((s for s in xl.sheet_names if "pmf" in s.lower()), None)
        if not sheet_name:
            st.error(f"❌ No sheet named like 'PMF' found. Sheets: {xl.sheet_names}")
            return None
        
        df = pd.read_excel(file_like_object, sheet_name=sheet_name, dtype=str)
        df.columns = [str(c).strip().upper() for c in df.columns]

        geo_col = next((c for c in df.columns if "GEO" in c), None)
        if not geo_col:
            st.error("❌ Geography column not found in PMF.")
            return None
        df.rename(columns={geo_col: "GEOGRAPHY"}, inplace=True)

        possible_season = ["SEASON", "PERIOD MAPPING", "PERIOD_MAPPING", "PERIOD_DEFINITION", "TIME_PERIODS"]
        season_col = next((c for c in df.columns if c in possible_season), None)
        if not season_col:
            st.error("❌ Season column not found in PMF.")
            return None
        df.rename(columns={season_col: "SEASON"}, inplace=True)

        df["GEOGRAPHY"] = df["GEOGRAPHY"].astype(str).str.upper().str.strip()
        df["SEASON"] = df["SEASON"].astype(str).str.upper().str.strip()
        return df
    except Exception as e:
        st.error(f"Error loading PMF: {e}")
        return None

@st.cache_data
def load_granular(file_like_object):
    try:
        xl = pd.ExcelFile(file_like_object)
        map_sheet = next((s for s in xl.sheet_names if "map" in s.lower()), None)
        if not map_sheet:
            st.error("❌ No 'MAP' sheet found in Granular file.")
            return None

        map_df = pd.read_excel(file_like_object, sheet_name=map_sheet, dtype=str)
        map_df.columns = [str(c).strip().upper() for c in map_df.columns]
        
        all_sheets = {s: pd.read_excel(file_like_object, sheet_name=s, dtype=str) for s in xl.sheet_names}
        return map_df, all_sheets
    except Exception as e:
        st.error(f"Error loading Granular Spec: {e}")
        return None

# =========================
# PROCESSING HELPERS
# =========================
def prepare_pmf_multipliers(pmf_df):
    pmf_vars = [c for c in pmf_df.columns if "_PMF" in c]
    pmf_long = pmf_df.melt(id_vars=["GEOGRAPHY", "SEASON"], value_vars=pmf_vars,
                          var_name="VARIABLE_PMF", value_name="MULTIPLIER")

    pmf_long["VARIABLE"] = pmf_long["VARIABLE_PMF"].str.replace("_PMF", "", regex=False)
    pmf_long["MULTIPLIER"] = pd.to_numeric(pmf_long["MULTIPLIER"], errors="coerce")

    pmf_dict = {(r.GEOGRAPHY, r.SEASON, r.VARIABLE): r.MULTIPLIER for r in pmf_long.itertuples()}
    return pmf_dict

def normalize_geo(name: str):
    return str(name).strip().upper().replace(".", "").replace("_", "").replace(" ", "")

def find_multiplier_for(sheet_name, season, var, mapcode_to_geo, pmf_dict_u):
    sheet_up = normalize_geo(sheet_name)
    season_up = str(season).strip().upper()
    var_up = str(var).strip().upper()

    geo = mapcode_to_geo.get(sheet_up)
    if geo:
        m = pmf_dict_u.get((normalize_geo(geo), season_up, var_up))
        if m is not None and not pd.isna(m):
            return m, geo

    m = pmf_dict_u.get((sheet_up, season_up, var_up))
    if m is not None and not pd.isna(m):
        return m, sheet_up

    return None, None

def apply_multipliers(granular_sheets, map_df, pmf_dict, selected_vars, tolerance_map, var_to_type, hard_limit, warning_limit):
    """
    Applying multipliers with Dynamic Tolerance Check, Hard Limits, Warnings, and Override Tab Logic.
    """
    map_df_copy = map_df.copy()
    map_df_copy["GEOGRAPHY"] = map_df_copy["GEOGRAPHY"].astype(str).str.strip().str.upper()
    map_df_copy["MAP"] = map_df_copy["MAP"].astype(str).str.strip().str.upper()
    mapcode_to_geo = dict(zip(map_df_copy["MAP"].tolist(), map_df_copy["GEOGRAPHY"].tolist()))

    pmf_dict_u = {
        (normalize_geo(k[0]), str(k[1]).strip().upper(), str(k[2]).strip().upper()): v
        for k, v in pmf_dict.items()
    }

    selected_vars_set = set([v.strip().upper() for v in selected_vars])
    multiplied_records, skipped_records = [], []
    updated_sheets = {}
    
    # 1. Identify the Override sheet (assumes the sheet name contains "override")
    override_sheet_name = next((s for s in granular_sheets.keys() if "override" in s.lower()), None)
    overridden_combinations = set()

    # Helper function to prevent code duplication between Override and Main tabs
    def process_row(df, idx, sheet_name_for_log, geo_for_lookup, var_col, contrib_col, min_col, max_col, is_override=False):
        raw_var = df.at[idx, var_col]
        raw_season = df.at[idx, contrib_col]

        if pd.isna(raw_var) or str(raw_var).strip() == "":
            return

        var = str(raw_var).strip().upper()
        season = str(raw_season).strip().upper() if not pd.isna(raw_season) else ""
        geo_norm = normalize_geo(geo_for_lookup)

        # If in override tab, log this specific (Geo, Var) combo so we skip it later in main sheets
        if is_override:
            overridden_combinations.add((geo_norm, var))

        # Check if it was overridden (only applies when processing main sheets)
        if not is_override and (geo_norm, var) in overridden_combinations:
            skipped_records.append([sheet_name_for_log, var, season, "SKIPPED_DUE_TO_OVERRIDE", None, None, None, None])
            return

        if var not in selected_vars_set:
            skipped_records.append([sheet_name_for_log, var, season, "VAR_NOT_SELECTED", None, None, None, None])
            return
            
        if season == "" or season in ["NAN", "NONE"]:
            skipped_records.append([sheet_name_for_log, var, season, "NO_SEASON", None, None, None, None])
            return

        multiplier, used_geo = find_multiplier_for(geo_for_lookup, season, var, mapcode_to_geo, pmf_dict_u)
        if multiplier is None or pd.isna(multiplier):
            skipped_records.append([sheet_name_for_log, var, season, "NO_MULTIPLIER_FOUND", None, None, None, None])
            return

        mult_val = float(multiplier)

        # --- HARD LIMIT CHECK ---
        if mult_val > hard_limit:
            skipped_records.append([sheet_name_for_log, var, season, f"EXCEEDS_HARD_LIMIT ({mult_val:.2f} > {hard_limit})", None, None, None, None])
            return

        # --- DYNAMIC TOLERANCE CHECK ---
        v_type = var_to_type.get(var)
        t_min, t_max = tolerance_map.get(v_type, (0.95, 1.05))

        if t_min <= mult_val <= t_max:
            skipped_records.append([sheet_name_for_log, var, season, f"SKIPPED_TOLERANCE ({mult_val:.3f})", None, None, None, None])
            return

        # --- APPLY MULTIPLIER ---
        try:
            old_min = float(df.at[idx, min_col]) if not pd.isna(df.at[idx, min_col]) else None
            old_max = float(df.at[idx, max_col]) if not pd.isna(df.at[idx, max_col]) else None
        except:
            skipped_records.append([sheet_name_for_log, var, season, "MIN_MAX_NOT_NUMERIC", None, None, None, None])
            return

        if old_min is None or old_max is None:
            return

        new_min = round(old_min * mult_val, 6)
        new_max = round(old_max * mult_val, 6)

        df.at[idx, min_col] = new_min
        df.at[idx, max_col] = new_max
        
        flag = "RED FLAG" if mult_val > warning_limit else ""
        multiplied_records.append([sheet_name_for_log, var, season, used_geo, multiplier, old_min, old_max, new_min, new_max, flag])

    # 2. Process Override Sheet First
    if override_sheet_name:
        df_over = granular_sheets[override_sheet_name].copy()
        cols_upper = [str(c).upper() for c in df_over.columns]
        
        # Override tab needs to have the Geography column to function
        if {"GEOGRAPHY", "VARIABLE", "CONTRIBUTION", "MIN", "MAX"}.issubset(cols_upper):
            geo_col = df_over.columns[cols_upper.index("GEOGRAPHY")]
            var_col = df_over.columns[cols_upper.index("VARIABLE")]
            contrib_col = df_over.columns[cols_upper.index("CONTRIBUTION")]
            min_col = df_over.columns[cols_upper.index("MIN")]
            max_col = df_over.columns[cols_upper.index("MAX")]

            for colname in [min_col, max_col]:
                df_over[colname] = (
                    df_over[colname]
                    .astype(str)
                    .str.replace("%", "", regex=False)
                    .str.replace(",", "", regex=False)
                    .apply(lambda x: pd.to_numeric(x, errors="coerce") if pd.notna(x) else np.nan)
                )

            for idx in df_over.index:
                raw_geo = df_over.at[idx, geo_col]
                if pd.isna(raw_geo) or str(raw_geo).strip() == "":
                    continue
                
                # Use raw_geo as the lookup geography, passing is_override=True
                process_row(df_over, idx, override_sheet_name, raw_geo, var_col, contrib_col, min_col, max_col, is_override=True)
                
        updated_sheets[override_sheet_name] = df_over

    # 3. Process Main Sheets
    for sheet_name, df_full in granular_sheets.items():
        if sheet_name == override_sheet_name:
            continue # Already processed
            
        df = df_full.copy()
        cols_upper = [str(c).upper() for c in df.columns]

        if not {"VARIABLE", "CONTRIBUTION", "MIN", "MAX"}.issubset(cols_upper):
            updated_sheets[sheet_name] = df
            continue

        var_col = df.columns[cols_upper.index("VARIABLE")]
        contrib_col = df.columns[cols_upper.index("CONTRIBUTION")]
        min_col = df.columns[cols_upper.index("MIN")]
        max_col = df.columns[cols_upper.index("MAX")]

        for colname in [min_col, max_col]:
            df[colname] = (
                df[colname]
                .astype(str)
                .str.replace("%", "", regex=False)
                .str.replace(",", "", regex=False)
                .apply(lambda x: pd.to_numeric(x, errors="coerce") if pd.notna(x) else np.nan)
            )

        for idx in df.index:
            # Use the sheet_name as the lookup geography
            process_row(df, idx, sheet_name, sheet_name, var_col, contrib_col, min_col, max_col, is_override=False)

        updated_sheets[sheet_name] = df

    multiplied_df = pd.DataFrame(
        multiplied_records,
        columns=["MAP_SHEET", "VARIABLE", "SEASON", "PMF_GEOGRAPHY_USED",
                 "MULTIPLIER", "OLD_MIN", "OLD_MAX", "NEW_MIN", "NEW_MAX", "FLAG"]
    )

    skipped_df = pd.DataFrame(
        skipped_records,
        columns=["MAP_SHEET", "VARIABLE", "SEASON", "REASON",
                 "OLD_MIN", "OLD_MAX", "NEW_MIN", "NEW_MAX"]
    )
    
    return updated_sheets, multiplied_df, skipped_df
# =========================
# OUTPUT HELPERS
# =========================
def create_output_excel(granular_sheets, updated_sheets):
    output_buffer = io.BytesIO()

    with pd.ExcelWriter(output_buffer, engine="openpyxl") as writer:
        for sheet_name, df in granular_sheets.items():
            df_to_save = updated_sheets.get(sheet_name, df)
            df_to_save.to_excel(writer, sheet_name=sheet_name, index=False)
    
    output_buffer.seek(0)
    wb = load_workbook(output_buffer)
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if ws.max_row <= 1:
             continue
             
        header = [cell.value for cell in ws[1] if cell.value]

        min_idx = max_idx = None
        for i, h in enumerate(header, start=1):
            if str(h).strip().upper() == "MIN":
                min_idx = i
            elif str(h).strip().upper() == "MAX":
                max_idx = i
        
        if min_idx is None and max_idx is None:
            continue

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            if min_idx:
                c = row[min_idx - 1]
                if isinstance(c.value, (int, float)):
                    c.number_format = numbers.FORMAT_PERCENTAGE_00
            if max_idx:
                c = row[max_idx - 1]
                if isinstance(c.value, (int, float)):
                    c.number_format = numbers.FORMAT_PERCENTAGE_00

    final_buffer = io.BytesIO()
    wb.save(final_buffer)
    wb.close()
    
    return final_buffer.getvalue()

def create_log_excel(multiplied_df, skipped_df):
    log_buffer = io.BytesIO()
    
    mult_summary = (
        multiplied_df.groupby("MAP_SHEET").size().reset_index(name="MULTIPLIED_COUNT")
        if not multiplied_df.empty else pd.DataFrame(columns=["MAP_SHEET","MULTIPLIED_COUNT"])
    )
    skip_summary = (
        skipped_df.groupby(["MAP_SHEET","REASON"]).size().reset_index(name="SKIPPED_COUNT")
        if not skipped_df.empty else pd.DataFrame(columns=["MAP_SHEET","REASON","SKIPPED_COUNT"])
    )
    
    summary_df = mult_summary.merge(skip_summary, on="MAP_SHEET", how="outer").fillna(0)
    
    if not summary_df.empty:
        summary_df.loc["Total"] = {
            "MAP_SHEET": "GRAND TOTAL",
            "MULTIPLIED_COUNT": mult_summary["MULTIPLIED_COUNT"].sum() if not mult_summary.empty else 0,
            "SKIPPED_COUNT": skip_summary["SKIPPED_COUNT"].sum() if not skip_summary.empty else 0
        }
    
    with pd.ExcelWriter(log_buffer, engine="openpyxl") as writer:
        multiplied_df.to_excel(writer, sheet_name="Multiplied", index=False)
        skipped_df.to_excel(writer, sheet_name="Skipped", index=False)
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        
        # --- APPLY RED FLAG FORMATTING ---
        ws = writer.sheets["Multiplied"]
        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid") # Light red background
        red_font = Font(color="9C0006") # Dark red text
        
        # Find index of the FLAG column
        flag_col_idx = None
        for col_idx, cell in enumerate(ws[1], 1):
            if cell.value == "FLAG":
                flag_col_idx = col_idx
                break
                
        if flag_col_idx:
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                if row[flag_col_idx - 1].value == "RED FLAG":
                    for cell in row:
                        cell.fill = red_fill
                        cell.font = red_font
    
    return log_buffer.getvalue(), summary_df

# =========================
# STREAMLIT UI
# =========================

defaults = {
    "step1_complete": False,
    "step2_complete": False,
    "step3_complete": False,
    "main_spec": None,
    "pmf": None,
    "map_df": None,
    "granular_sheets": None,
    "selected_types": None,
    "selected_vars": None,
    "pmf_dict": None,
    "updated_sheets": None,
    "multiplied_df": None,
    "skipped_df": None,
    "output_file_bytes": None,
    "log_file_bytes": None,
    "log_summary_df": None,
    "tolerance_map": {}, 
    "var_to_type": {},
    "hard_limit": 100.0,
    "warning_limit": 50.0
}
for key, value in defaults.items():
    if key not in st.session_state:
        st.session_state[key] = value

tab1, tab2, tab3, tab4 = st.tabs([
    "Step 1: Upload Files", 
    "Step 2: Select Types & Tolerance", 
    "Step 3: Run Process", 
    "Step 4: Download Results"
])

# =========================
# TAB 1: FILE UPLOAD
# =========================
with tab1:
    st.header("Step 1: Upload Files")
    st.info("Please upload the three required Excel files. Processing will begin automatically.")

    gran_file = st.file_uploader("1. Select Granular Spec (Excel)", type=['xlsx'])
    pmf_file = st.file_uploader("2. Select PMF File (Excel)", type=['xlsx'])
    main_file = st.file_uploader("3. Select Main Spec (Excel)", type=['xlsx'])

    if gran_file and pmf_file and main_file:
        with st.spinner("Loading files..."):
            main_spec_df = load_main_spec(main_file)
            pmf_df = load_pmf(pmf_file)
            map_df, granular_sheets_dict = load_granular(gran_file)

        if main_spec_df is not None and pmf_df is not None and map_df is not None:
            st.session_state["main_spec"] = main_spec_df
            st.session_state["pmf"] = pmf_df
            st.session_state["map_df"] = map_df
            st.session_state["granular_sheets"] = granular_sheets_dict
            st.session_state["step1_complete"] = True
            
            st.session_state["var_to_type"] = dict(zip(main_spec_df["VARIABLE"], main_spec_df["TYPE"]))
                
            st.success("✅ All files loaded and validated successfully!")
            st.subheader("File Previews (First 5 Rows)")
            st.write("**Main Spec (Processed)**")
            st.dataframe(main_spec_df.head())
            st.write("**PMF (Processed)**")
            st.dataframe(pmf_df.head())
            st.write("**Granular MAP (Processed)**")
            st.dataframe(map_df.head())
        else:
            st.error("One or more files failed to load. Please check errors above.")
            st.session_state["step1_complete"] = False

# =========================
# TAB 2: TYPE & TOLERANCE SELECTION
# =========================
with tab2:
    st.header("Step 2: Configuration")
    if not st.session_state["step1_complete"]:
        st.warning("Please upload all files in Step 1 first.")
    else:
        # --- VARIABLE TYPES ---
        st.subheader("1. Variable Types")
        
        choice_map = {
            "Base": ["Base"],
            "Incremental": ["Incremental"],
            "Base and Incremental": ["Base", "Incremental"]
        }
        
        type_choice = st.radio(
            "Choose variable types to process:",
            options=choice_map.keys(),
            index=2
        )
        
        selected_types = choice_map[type_choice]
        selected_vars = st.session_state["main_spec"][
            st.session_state["main_spec"]["TYPE"].isin(selected_types)
        ]["VARIABLE"].tolist()
        
        st.session_state["selected_types"] = selected_types
        st.session_state["selected_vars"] = selected_vars
        
        st.divider()

        # --- DYNAMIC TOLERANCE SETTINGS ---
        st.subheader("2. Tolerance Level")
        st.info("Set the skip tolerance range for **each** selected variable type.")
        
        tolerance_map_temp = {}
        
        with st.container():
            for t_type in selected_types:
                st.markdown(f"**{t_type}** Tolerance Settings")
                c1, c2 = st.columns(2)
                
                t_min = c1.number_input(f"Min ({t_type})", value=0.95, step=0.01, format="%.2f", key=f"min_{t_type}")
                t_max = c2.number_input(f"Max ({t_type})", value=1.05, step=0.01, format="%.2f", key=f"max_{t_type}")
                
                tolerance_map_temp[t_type] = (t_min, t_max)
                st.write("") 
        
        st.session_state["tolerance_map"] = tolerance_map_temp

        st.divider()
        
        # --- MULTIPLIER LIMITS ---
        st.subheader("3. Multiplier Limits")
        st.info("Set the hard limit (skips row) and warning threshold (flags red in log).")
        c3, c4 = st.columns(2)
        st.session_state["warning_limit"] = c3.number_input("Warning Threshold (> flags red)", value=50.0, step=1.0)
        st.session_state["hard_limit"] = c4.number_input("Hard Limit (> skips processing)", value=100.0, step=1.0)

        st.divider()
        st.session_state["step2_complete"] = True

        st.success(f"✅ Configuration Ready: **{len(selected_vars)} variables** selected across **{len(selected_types)} types**.")
        
        with st.expander("Click to see all selected variables"):
            st.dataframe(selected_vars)

# =========================
# TAB 3: RUN PROCESS
# =========================
with tab3:
    st.header("Step 3: Apply Multipliers")
    if not st.session_state["step2_complete"]:
        st.warning("Please complete Steps 1 and 2 first.")
    else:
        st.info("This step will prepare the PMF multipliers and apply them to all sheets.")
        
        if st.button("🚀 Run Multiplier Process", type="primary", use_container_width=True):
            
            # --- BOSE ANIMATION INJECTION ---
            loader_placeholder = st.empty()
            loader_placeholder.markdown("""
            <div class="bose-loader-overlay">
                <div class="bose-juggler">
                    <span class="bose-letter">B</span>
                    <span class="bose-letter">O</span>
                    <span class="bose-letter">S</span>
                    <span class="bose-letter">E</span>
                </div>
                <div class="bose-status">Processing Data...</div>
            </div>
            """, unsafe_allow_html=True)
            
            time.sleep(1.2)
            
            try:
                # Step 4: Prepare PMF
                pmf_dict = prepare_pmf_multipliers(st.session_state["pmf"])
                st.session_state["pmf_dict"] = pmf_dict
                
                # Step 5: Apply Multipliers (With DYNAMIC Tolerance & Limits)
                updated_sheets, multiplied_df, skipped_df = apply_multipliers(
                    st.session_state["granular_sheets"],
                    st.session_state["map_df"],
                    st.session_state["pmf_dict"],
                    st.session_state["selected_vars"],
                    st.session_state["tolerance_map"], 
                    st.session_state["var_to_type"],
                    st.session_state["hard_limit"],
                    st.session_state["warning_limit"]
                )
                
                st.session_state["updated_sheets"] = updated_sheets
                st.session_state["multiplied_df"] = multiplied_df
                st.session_state["skipped_df"] = skipped_df
                st.session_state["step3_complete"] = True
                
                loader_placeholder.empty() # Remove Animation
                st.success("✅ Multipliers applied successfully!")
                st.metric("Records Multiplied", len(multiplied_df))
                st.metric("Records Skipped", len(skipped_df))
                
            except Exception as e:
                loader_placeholder.empty() # Remove Animation
                st.error(f"An error occurred during processing: {e}")
                st.session_state["step3_complete"] = False

# =========================
# TAB 4: DOWNLOAD RESULTS
# =========================
with tab4:
    st.header("Step 4: Download Results")
    if not st.session_state["step3_complete"]:
        st.warning("Please run the process in Step 3 first.")
    else:
        st.success("🎉 Process completed! Your files are ready for download.")
        
        if st.session_state["output_file_bytes"] is None:
            with st.spinner("Generating Excel file..."):
                output_bytes = create_output_excel(
                    st.session_state["granular_sheets"],
                    st.session_state["updated_sheets"]
                )
            st.session_state["output_file_bytes"] = output_bytes
        
        today_str = date.today().isoformat()
        types_str = '_'.join(st.session_state['selected_types'])
        granular_outfile_name = f"Granular_Updated_{types_str}_{today_str}.xlsx"
        log_outfile_name = f"Granular_Log_{types_str}_{today_str}.xlsx"

        st.download_button(
            label="💾 Download Updated Granular File",
            data=st.session_state["output_file_bytes"],
            file_name=granular_outfile_name,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True
        )
        
        st.markdown("---")
        
        generate_log = st.checkbox("Generate detailed log file?", value=True)
        
        if generate_log:
            if st.session_state["log_file_bytes"] is None:
                log_bytes, summary_df = create_log_excel(
                    st.session_state["multiplied_df"],
                    st.session_state["skipped_df"]
                )
                st.session_state["log_file_bytes"] = log_bytes
                st.session_state["log_summary_df"] = summary_df

            if st.session_state["log_file_bytes"]:
                st.download_button(
                    label="📋 Download Log File",
                    data=st.session_state["log_file_bytes"],
                    file_name=log_outfile_name,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )
                
                st.subheader("Log Summary")
                st.dataframe(st.session_state["log_summary_df"])

        with st.expander("View Multiplied Records Details"):
            st.dataframe(st.session_state["multiplied_df"])
        
        with st.expander("View Skipped Records Details"):
            st.dataframe(st.session_state["skipped_df"])

    st.markdown("---")
    if st.button("🔄 Start Over", type="secondary", use_container_width=True):
        st.session_state.clear()
        st.rerun()